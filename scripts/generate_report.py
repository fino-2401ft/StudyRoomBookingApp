from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_run_font(run, size=12, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading_custom(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=14, bold=True)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_para(doc, text, bold=False, italic=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    run.italic = italic
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run)
    p.paragraph_format.space_after = Pt(3)
    return p


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=11, bold=bold)


def build_report() -> Document:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MINI-PROJECT SHORT TECHNICAL REPORT")
    set_run_font(r, size=16, bold=True)

    meta_lines = [
        ("Course: ", "Cross-Platform Mobile App Development (VKU)"),
        (
            "Mini-Project Title: ",
            "Mini-Project 2 — Study Room Booking App (React Native & Expo)",
        ),
        ("Team / Student Name: ", "Nguyễn Lê Quốc Anh"),
        ("Submission Date: ", "08/10/2026"),
    ]
    for label, value in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(label)
        set_run_font(r1, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2)
        p.paragraph_format.space_after = Pt(2)

    doc.add_paragraph()

    # Section 1
    add_heading_custom(doc, "1. GENERAL INFORMATION & DELIVERABLE LINKS")
    add_para(doc, "Team Members:", bold=True)
    add_bullet(
        doc,
        "Nguyễn Lê Quốc Anh — Student ID: 22ITxxx — Role: Full-stack Mobile Developer "
        "(UI, State, Booking Logic) — Contribution: 100%",
    )
    add_para(
        doc,
        "Live Demo URL: Expo Go (LAN/Tunnel) — run `npx expo start` in the project root; "
        "open with Expo Go on a physical device.",
    )
    add_para(doc, "GitHub Repository: [Update with your repository URL]")
    add_para(
        doc,
        "Video Demo (Optional): [Update with YouTube / Drive link if available]",
    )

    # Section 2
    add_heading_custom(doc, "2. FEATURE IMPLEMENTATION CHECKLIST")

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = [
        "#",
        "Required Feature",
        "Status",
        "Implementation Details & Acceptance Level",
    ]
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, center=True)
        shade_cell(table.rows[0].cells[i], "E7C9D0")

    rows = [
        (
            "1",
            "Room Discovery with high-performance FlatList",
            "Complete",
            "RoomsScreen renders campus rooms via FlatList (initialNumToRender=6, "
            "maxToRenderPerBatch=8, removeClippedSubviews). RoomCard is React.memo with "
            "custom areEqual. Shows image, building/floor, capacity badge, Available Now / Occupied.",
        ),
        (
            "2",
            "Multi-parameter search & filters",
            "Complete",
            "FilterBar filters by building (A/B/C/V), capacity bands (2–8, 9–14, 15–20), "
            "equipment (Projector, Whiteboard, High-end PCs, Air Conditioning), text search, "
            "and Available only. Filter state is stored in Zustand.",
        ),
        (
            "3",
            "Interactive 7-day date picker & 2-hour time slots",
            "Complete",
            "DatePicker shows the next 7 days. TimeSlotPicker offers fixed slots: "
            "07:30–09:30, 09:30–11:30, 13:00–15:00, 15:00–17:00.",
        ),
        (
            "4",
            "Time-slot conflict prevention",
            "Complete",
            "Booked and past slots are disabled visually. createBooking double-checks "
            "isSlotBooked before commit and returns an error on conflict. Seeded demo "
            "bookings illustrate locked slots.",
        ),
        (
            "5",
            "Unique booking code + QR check-in modal",
            "Complete",
            "generateBookingCode() produces VKU-XXXXXX. BookingSuccessModal shows an "
            "interactive QR (react-native-qrcode-svg) with booking payload for check-in.",
        ),
        (
            "6",
            "Global state with Zustand (useBookingStore)",
            "Complete",
            "Manages user session, active bookings, filters, create/cancel booking. "
            "Selectors are used across Rooms, Booking, My Bookings, and Profile screens.",
        ),
        (
            "7",
            "Local persistence via AsyncStorage",
            "Complete",
            "zustand/middleware persist + createJSONStorage(() => AsyncStorage) under key "
            "vku-booking-store. Session and bookings survive app restarts; App.tsx waits "
            "for hydration before rendering.",
        ),
        (
            "8",
            "Local notifications (15 minutes before slot)",
            "Complete",
            "expo-notifications schedules a DATE trigger 15 minutes before slot start. "
            "Cancel booking also cancels the reminder. Profile can request notification permission.",
        ),
    ]

    for row in rows:
        cells = table.add_row().cells
        set_cell_text(cells[0], row[0], center=True)
        set_cell_text(cells[1], row[1])
        set_cell_text(cells[2], row[2], center=True)
        set_cell_text(cells[3], row[3])

    # Section 3
    add_heading_custom(doc, "3. TECHNICAL ARCHITECTURE & PROJECT STRUCTURE")
    add_para(
        doc,
        "Stack: React Native 0.86 + Expo SDK 57 + TypeScript, React Navigation "
        "(native stack + bottom tabs), Zustand, AsyncStorage, expo-notifications, "
        "react-native-qrcode-svg.",
    )
    add_para(doc, "Directory structure:", bold=True)
    add_bullet(
        doc,
        "src/components/ — RoomCard, FilterBar, DatePicker, TimeSlotPicker, "
        "BookingSuccessModal, CrystalBackground, FallingPetals",
    )
    add_bullet(doc, "src/data/ — Mock rooms (12 VKU rooms) and TIME_SLOTS definitions")
    add_bullet(
        doc,
        "src/navigation/ — AppNavigator (tabs: Rooms / Bookings / Profile; stack: RoomDetail, Booking)",
    )
    add_bullet(
        doc,
        "src/screens/ — RoomsScreen, RoomDetailScreen, BookingScreen, MyBookingsScreen, ProfileScreen",
    )
    add_bullet(doc, "src/store/ — useBookingStore (Zustand + persist)")
    add_bullet(
        doc,
        "src/theme/, src/types/, src/utils/ — crystal-pink theme, shared types, "
        "date/booking/notification helpers",
    )

    add_para(doc, "State management flow:", bold=True)
    add_bullet(
        doc,
        "UI actions call store methods (setFilters, createBooking, cancelBooking).",
    )
    add_bullet(
        doc,
        "createBooking validates conflicts → generates code → schedules notification → "
        "appends booking → AsyncStorage persist.",
    )
    add_bullet(
        doc,
        "Screens subscribe with fine-grained selectors to limit re-renders "
        "(supports the FlatList 60fps goal).",
    )

    add_para(doc, "Exception / edge handling:", bold=True)
    add_bullet(
        doc,
        "Conflict: booked slots disabled; createBooking returns { ok: false, error } and shows Alert.",
    )
    add_bullet(doc, "Past slots blocked via isSlotInPast.")
    add_bullet(
        doc,
        "Notification permission denied: booking still succeeds; reminder is skipped gracefully.",
    )
    add_bullet(
        doc,
        "Hydration: App waits for Zustand rehydration before rendering the navigator.",
    )
    add_bullet(
        doc,
        "LAN/device networking: Expo host forced to Wi-Fi IP (avoid VMware adapters) "
        "or --tunnel when Expo Go cannot connect.",
    )

    # Section 4
    add_heading_custom(doc, "4. EMPIRICAL EVIDENCE & SCREENSHOTS")
    add_para(
        doc,
        "Insert 3–4 annotated screenshots of the application running on Expo Go / emulator. "
        "Suggested captures:",
        italic=True,
    )
    add_bullet(
        doc,
        "Screenshot 1 — Rooms tab: FlatList of rooms with crystal-pink UI, falling-petal "
        "background, multi-parameter FilterBar, Available Now / Occupied badges.",
    )
    add_bullet(
        doc,
        "Screenshot 2 — Booking screen: 7-day DatePicker + TimeSlotPicker with booked "
        "slots visually disabled (conflict prevention).",
    )
    add_bullet(
        doc,
        "Screenshot 3 — Booking success modal: unique VKU-XXXXXX code and QR check-in.",
    )
    add_bullet(
        doc,
        "Screenshot 4 — My Bookings / Profile: active reservations, cancel flow, "
        "AsyncStorage-backed session, notification permission action.",
    )
    add_para(
        doc,
        "[Paste screenshots below this line in the Word document before submission.]",
        italic=True,
    )

    # Section 5
    add_heading_custom(doc, "5. TECHNICAL CHALLENGES & RESOLUTIONS")
    add_para(
        doc,
        "Challenge 1 — Expo Go / Safari cannot connect to Metro on iPhone",
        bold=True,
    )
    add_para(
        doc,
        "On Windows with VMware adapters, Expo often advertised a virtual NIC IP "
        "(e.g. 192.168.140.x) instead of the Wi-Fi address (192.168.100.54). Scanning the "
        "QR opened an unreachable URL in Safari (“cannot connect to the server”).",
    )
    add_para(
        doc,
        "Resolution: Restart Expo with REACT_NATIVE_PACKAGER_HOSTNAME set to the Wi-Fi "
        "IP and --lan mode; open via Expo Go (not Safari). Fallback: npx expo start --tunnel "
        "when LAN is blocked by firewall/subnet.",
    )

    add_para(
        doc,
        "Challenge 2 — Balancing tab-bar safe area with design requirements",
        bold=True,
    )
    add_para(
        doc,
        "The floating tab bar was initially lifted too high (+18px extra). Feedback "
        "required it near the bottom edge while remaining clear of the iPhone home indicator.",
    )
    add_para(
        doc,
        "Resolution: Set bottom edgeGap = 6px and compute paddingBottom from "
        "useSafeAreaInsets().bottom so icons stay above the home indicator without floating "
        "mid-screen. Content lists use paddingBottom ≈ 100 to avoid overlap.",
    )

    add_para(doc, "Challenge 3 — Smooth list performance with rich UI", bold=True)
    add_para(
        doc,
        "Crystal background + falling petals + image cards risked jank on FlatList scroll.",
    )
    add_para(
        doc,
        "Resolution: Memoized RoomCard, stable onPress via useCallback, FlatList windowing "
        "props, and petals rendered once in CrystalBackground with pointerEvents=\"none\" "
        "and useNativeDriver animations.",
    )

    doc.add_paragraph()
    note = doc.add_paragraph()
    r = note.add_run(
        "Note: Replace Student ID, GitHub URL, Live Demo / Video links, and insert real "
        "device screenshots before final submission."
    )
    set_run_font(r, size=11)
    r.italic = True

    return doc


def main():
    doc = build_report()
    out_paths = [
        Path(
            r"c:\Users\HP VICTUS\Documents\VKU_Subjects\HK7\Multi_Platform\StudyRoomBookingApp\Mini-Project-2-StudyRoomBooking-Report.docx"
        ),
        Path(
            r"C:\Users\HP VICTUS\Documents\VKU_Subjects\HK7\CloudComputing\NguyenLeQuocAnh\MidTerm\Mini-Project-2-StudyRoomBooking-Report.docx"
        ),
    ]
    for out in out_paths:
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(out)
        print("Saved:", out)


if __name__ == "__main__":
    main()
